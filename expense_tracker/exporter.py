"""Export expenses and income: CSV for any program, a formatted Excel workbook, or the daily table."""

from __future__ import annotations

import csv
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from . import APP_NAME, __version__
from .daily import DAY_FILL, DailyTable, amount_text, header_color
from .daily import TOTAL_FILL as DAILY_TOTAL_FILL
from .storage import SYMBOLS, Expense, Income, format_rate, total_in


@dataclass(frozen=True)
class Column:
    title: str
    value: Callable[[Expense | Income], str | date | Decimal | None]
    kind: str = "text"  # "text", "date", "money" or "rate"
    currency: str | None = None  # money columns: always this currency, or None for the record's own


_MONEY_COLUMNS = [
    Column("Currency", lambda r: r.currency),
    Column("Amount", lambda r: r.amount, "money"),
    Column("Rate (PEN per USD)", lambda r: r.rate, "rate"),
    Column("Amount in PEN", lambda r: r.amount_pen, "money", "PEN"),
    Column("Amount in USD", lambda r: r.amount_usd, "money", "USD"),
]
EXPENSE_COLUMNS = [
    Column("Date", lambda e: e.date, "date"),
    Column("Category", lambda e: e.category),
    Column("Subcategory", lambda e: e.subcategory),
    Column("Description", lambda e: e.description),
    Column("Payment method", lambda e: e.payment_method),
    Column("Tags", lambda e: ", ".join(e.tags)),
    *_MONEY_COLUMNS,
]
INCOME_COLUMNS = [
    Column("Date", lambda i: i.date, "date"),
    Column("Source", lambda i: i.source),
    Column("Description", lambda i: i.description),
    *_MONEY_COLUMNS,
]


def _totalled(columns: list[Column]) -> list[int]:
    """Indexes of the columns that get a total: amounts in one fixed currency."""
    return [i for i, c in enumerate(columns) if c.kind == "money" and c.currency]


def _total_label(records: list) -> str:
    no_rate = sum(r.rate is None for r in records)
    return f"Total ({len(records)}{f'; {no_rate} without exchange rate' if no_rate else ''})"


# ---------- CSV ----------

def _csv_cell(column: Column, record: Expense | Income) -> str:
    value = column.value(record)
    if value is None:
        return ""
    if column.kind == "date":
        return value.isoformat()
    if column.kind == "money":
        return f"{value:.2f}"
    if column.kind == "rate":
        return format_rate(value)
    return str(value)


def _write_csv(records: list, columns: list[Column], path: Path) -> None:
    # utf-8-sig adds a BOM, so Excel reads accents and emoji correctly.
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([c.title for c in columns])
        for r in records:
            writer.writerow([_csv_cell(c, r) for c in columns])
        if records:
            totals = [""] * len(columns)
            totals[0] = _total_label(records)
            for i in _totalled(columns):
                totals[i] = f"{total_in(records, columns[i].currency)[0]:.2f}"
            writer.writerow([])  # a blank line keeps the totals out of Excel's sort and filter range
            writer.writerow(totals)


def export_expenses_csv(expenses: list[Expense], path: Path) -> None:
    _write_csv(expenses, EXPENSE_COLUMNS, path)


def export_income_csv(incomes: list[Income], path: Path) -> None:
    _write_csv(incomes, INCOME_COLUMNS, path)


# ---------- Excel ----------

HEADER_ROW = 4  # rows 1-2: title and notes
EXPENSE_FILL, INCOME_FILL = "E76F51", "43AA8B"  # the app's expense and income colors
MUTED, TOTAL_FILL = "6C757D", "F1F3F5"


def _excel_format(column: Column, record: Expense | Income) -> str | None:
    if column.kind == "date":
        return "d mmm yyyy"
    if column.kind == "money":
        return f'"{SYMBOLS[column.currency or record.currency]}" #,##0.00'
    if column.kind == "rate":
        return "0.00##"
    return None


def _shown_width(column: Column, record: Expense | Income, value) -> int:
    """About how many characters Excel shows for this value, to size the column."""
    if value is None:
        return 0
    if column.kind == "date":
        return 11  # "29 Sep 2026"
    if column.kind == "money":
        return len(f"{SYMBOLS[column.currency or record.currency]} {value:,.2f}")
    if column.kind == "rate":
        return len(format_rate(value))
    return len(str(value))


def _fill_sheet(ws, columns: list[Column], records: list, title: str, note: str, color: str, empty: str) -> None:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = note
    ws["A2"].font = Font(size=9, color=MUTED)
    ws.sheet_properties.tabColor = color
    for col, c in enumerate(columns, 1):
        cell = ws.cell(HEADER_ROW, col, c.title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="right" if c.kind in ("money", "rate") else "left")
    widths = [len(c.title) + 3 for c in columns]  # + room for the filter arrow
    first = HEADER_ROW + 1
    for row, r in enumerate(records, first):
        for col, c in enumerate(columns, 1):
            value = c.value(r)
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value) or None  # Excel files can't hold control characters
            cell = ws.cell(row, col, value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # text, not a formula
            fmt = _excel_format(c, r)
            if fmt:
                cell.number_format = fmt
            if c.kind == "date":
                cell.alignment = Alignment(horizontal="left")  # like its header, away from the next column
            widths[col - 1] = max(widths[col - 1], _shown_width(c, r, value))
    last = first + len(records) - 1
    end = get_column_letter(len(columns))
    if records:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{end}{last}"
        # SUBTOTAL ignores rows hidden by the filter, so the totals follow what's filtered in Excel.
        amount = get_column_letter(next(i for i, c in enumerate(columns, 1) if c.title == "Amount"))
        ws.cell(last + 1, 1, f'="Total ("&SUBTOTAL(103,{amount}{first}:{amount}{last})&")"')
        for i in _totalled(columns):
            letter = get_column_letter(i + 1)
            cell = ws.cell(last + 1, i + 1, f"=SUBTOTAL(109,{letter}{first}:{letter}{last})")
            cell.number_format = f'"{SYMBOLS[columns[i].currency]}" #,##0.00'
        line = Side(style="thin", color=MUTED)
        for col in range(1, len(columns) + 1):
            cell = ws.cell(last + 1, col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            cell.border = Border(top=line)
    else:
        ws.cell(first, 1, empty).font = Font(italic=True, color=MUTED)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 50)
    ws.freeze_panes = ws.cell(first, 1)
    # Printing: landscape, all columns on one page width, the header repeated on every page.
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def export_xlsx(
    path: Path, expenses: list[Expense], incomes: list[Income], period: str, user: str = "", filters: str = "",
) -> None:
    """A workbook with an Expenses and an Income sheet, formatted for reading, filtering and printing.

    `period` names what was exported (e.g. "September 2026"); `filters` describes whatever narrowed
    down the expenses (the Income sheet always has all income of the period).
    """
    from openpyxl import Workbook  # only needed here, so everything else works without openpyxl

    today = date.today()
    exported = f"Exported {today.day} {today:%b %Y} with {APP_NAME}"
    wb = Workbook()
    wb.properties.creator = f"{APP_NAME} {__version__}"
    expenses_sheet = wb.active
    expenses_sheet.title = "Expenses"
    _fill_sheet(expenses_sheet, EXPENSE_COLUMNS, expenses, f"Expenses — {period}",
                " · ".join(filter(None, [user, filters, exported])), EXPENSE_FILL, "No expenses")
    _fill_sheet(wb.create_sheet("Income"), INCOME_COLUMNS, incomes, f"Income — {period}",
                " · ".join(filter(None, [user, exported])), INCOME_FILL, "No income")
    wb.save(path)


# ---------- daily expenses table (Excel) ----------

DAILY_FIRST_ROW = 4  # rows 1-2: title and notes


def _outline(ws, top: int, left: int, bottom: int, right: int, side) -> None:
    """Draw `side` around a block of cells, keeping the borders the cells already have."""
    from openpyxl.styles import Border

    for row in range(top, bottom + 1):
        for col in range(left, right + 1):
            cell, old = ws.cell(row, col), ws.cell(row, col).border
            cell.border = Border(left=side if col == left else old.left, right=side if col == right else old.right,
                                 top=side if row == top else old.top, bottom=side if row == bottom else old.bottom)


def export_daily_xlsx(path: Path, table: DailyTable, title: str, colors: dict[str, str], note: str = "") -> None:
    """The daily expenses table as in the app: a block per day, each category with item, description and
    amount, and a yellow sums row (with formulas, so the sums follow any change made in Excel)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    today = date.today()
    wb = Workbook()
    wb.properties.creator = f"{APP_NAME} {__version__}"
    ws = wb.active
    ws.title = "Daily expenses"
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = " · ".join(filter(None, [note, f"Amounts in {table.currency}",
                                        f"Exported {today.day} {today:%b %Y} with {APP_NAME}"]))
    ws["A2"].font = Font(size=9, color=MUTED)

    money = f'"{SYMBOLS[table.currency]}" #,##0.00'
    day_fill = PatternFill("solid", fgColor=DAY_FILL[1:].upper())
    thick, thin = Side(style="medium", color="333333"), Side(style="thin", color="333333")
    bold, right = Font(bold=True), Alignment(horizontal="right")
    first_col = {cat: 2 + 3 * i for i, cat in enumerate(table.categories)}  # each category: item, description, amount
    total_col = 2 + 3 * len(table.categories)
    widths = {1: 19, total_col: 13}  # A: bold "2026-10-07 (Wed)"

    def width(col: int, text: str, cap: int = 40) -> None:
        widths[col] = min(cap, max(widths.get(col, 6), len(text) + 2))

    row = DAILY_FIRST_ROW
    for day in table.days:
        header, footer = row, row + max(day.rows, 1) + 1
        for r in (header, footer):
            cell = ws.cell(r, 1, day.day)
            cell.number_format, cell.font, cell.fill = "yyyy-mm-dd (ddd)", bold, day_fill
            cell.alignment = Alignment(horizontal="left")
        for col in range(2, total_col + 1):
            ws.cell(footer, col).fill = day_fill
        for cat, col in first_col.items():
            amount_col = col + 2
            cell = ws.cell(header, amount_col, cat)
            cell.font, cell.alignment = bold, right
            cell.fill = PatternFill("solid", fgColor=header_color(colors.get(cat, "#8d99ae"))[1:].upper())
            width(amount_col, cat)
            items = day.items.get(cat, [])
            for i, e in enumerate(items, start=1):
                ws.cell(header + i, col, e.subcategory or None)
                ws.cell(header + i, col + 1, e.description or None)
                value = e.amount_in(table.currency)
                cell = ws.cell(header + i, amount_col, value if value is not None else amount_text(e, table.currency))
                cell.number_format = money
                width(col, e.subcategory)
                width(col + 1, e.description)
                width(amount_col, amount_text(e, table.currency))
            if items:
                letter = get_column_letter(amount_col)
                cell = ws.cell(footer, amount_col, f"=SUM({letter}{header + 1}:{letter}{footer - 1})")
                cell.number_format, cell.font = money, bold
            _outline(ws, header, col, footer, amount_col, thick)
        cell = ws.cell(header, total_col, "Day total")
        cell.font, cell.alignment = Font(color=MUTED), right
        sums = ",".join(f"{get_column_letter(col + 2)}{footer}" for cat, col in first_col.items() if cat in day.items)
        cell = ws.cell(footer, total_col, f"=SUM({sums})")
        cell.number_format, cell.font = money, bold
        _outline(ws, header, 1, footer, 1, thick)
        _outline(ws, header, total_col, footer, total_col, thick)
        for col in range(1, total_col + 1):  # the line above the sums
            cell, old = ws.cell(footer, col), ws.cell(footer, col).border
            cell.border = Border(left=old.left, right=old.right, top=thin, bottom=old.bottom)
        row = footer + 1

    if table.days:
        # Totals: column A only holds dates on the days' first and sums rows, and the first row's
        # amount cells hold category names (text, which SUMIF skips): so this adds up the daily sums.
        last, row = row - 1, row + 1
        total_fill = PatternFill("solid", fgColor=DAILY_TOTAL_FILL[1:].upper())
        for col in range(1, total_col + 1):
            ws.cell(row, col).fill = total_fill
        ws.cell(row, 1, "Total").font = bold
        for col in (*(c + 2 for c in first_col.values()), total_col):
            letter = get_column_letter(col)
            cell = ws.cell(row, col, f'=SUMIF($A${DAILY_FIRST_ROW}:$A${last},"<>",{letter}{DAILY_FIRST_ROW}:{letter}{last})')
            cell.number_format, cell.font = money, bold
        _outline(ws, row, 1, row, total_col, thick)
    else:
        ws.cell(DAILY_FIRST_ROW, 1, "No expenses in this period").font = Font(italic=True, color=MUTED)

    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "B1"  # dates stay in view when scrolling sideways
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
