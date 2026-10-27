import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from expense_tracker.daily import amount_text, build_daily_table, header_color
from expense_tracker.exporter import export_daily_xlsx
from expense_tracker.storage import Expense

D = Decimal


def sample() -> list[Expense]:
    return [
        Expense(D("8.50"), "Food", "Espresso", date(2026, 10, 3), subcategory="Coffee", rate=D("3.75")),
        Expense(D("12.00"), "Transport", "To the market", date(2026, 10, 3), subcategory="Taxi", rate=D("3.75")),
        Expense(D("18.00"), "food", "Menu del día", date(2026, 10, 3), subcategory="Restaurant", rate=D("3.75")),
        Expense(D("10.00"), "Zoo", "Tickets", date(2026, 10, 4), currency="USD", rate=D("3.75")),
        Expense(D("5.00"), "Books", "Old entry", date(2026, 10, 4), currency="USD"),  # no exchange rate
        Expense(D("45.00"), "Food", "Ceviche", date(2026, 10, 4), subcategory="Restaurant", rate=D("3.75")),
        Expense(D("99.00"), "Food", "Last month", date(2026, 9, 30), rate=D("3.75")),
    ]


class DailyTableTest(unittest.TestCase):
    def test_table(self):
        t = build_daily_table(sample(), "2026-10", "PEN", ["Transport", "Food", "Housing"])
        # The user's category order first, then the others A→Z; "food" and "Food" are one category
        self.assertEqual(t.categories, ["Transport", "Food", "Books", "Zoo"])
        self.assertEqual([d.day for d in t.days], [date(2026, 10, 3), date(2026, 10, 4)])
        first = t.days[0]
        self.assertEqual([e.description for e in first.items["Food"]], ["Espresso", "Menu del día"])  # as entered
        self.assertEqual((first.rows, first.sums, first.total),
                         (2, {"Food": D("26.50"), "Transport": D("12.00")}, D("38.50")))
        second = t.days[1]
        self.assertEqual(second.sums, {"Zoo": D("37.50"), "Books": D("0"), "Food": D("45.00")})  # no rate: not counted
        self.assertEqual(t.sums, {"Transport": D("12.00"), "Food": D("71.50"), "Books": D("0"), "Zoo": D("37.50")})
        self.assertEqual((t.total, t.count, t.missing_rate), (D("121.00"), 6, 1))

    def test_periods(self):
        self.assertEqual(build_daily_table(sample(), "2026", "PEN", []).count, 7)
        self.assertEqual(build_daily_table(sample(), None, "PEN", []).count, 7)
        empty = build_daily_table(sample(), "2026-11", "PEN", [])
        self.assertEqual((empty.days, empty.categories, empty.total), ([], [], D("0")))

    def test_amount_text_and_colors(self):
        usd, old = sample()[3], sample()[4]
        self.assertEqual((amount_text(usd, "PEN"), amount_text(usd, "USD")), ("S/ 37.50", "US$ 10.00"))
        self.assertEqual(amount_text(old, "PEN"), "US$ 5.00 (no rate)")
        self.assertEqual(header_color("#e76f51"), "#f5c5b9")  # each channel 60% of the way to white
        self.assertEqual(header_color("#ffffff"), "#ffffff")


class DailyExcelTest(unittest.TestCase):
    def test_layout(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl isn't installed")
        table = build_daily_table(sample(), "2026-10", "PEN", ["Transport", "Food"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "daily.xlsx"
            export_daily_xlsx(path, table, "Daily expenses — October 2026", {"Food": "#e76f51"}, "Ana")
            ws = load_workbook(path).active
        self.assertEqual(ws["A1"].value, "Daily expenses — October 2026")
        self.assertTrue(ws["A2"].value.startswith("Ana · Amounts in PEN · Exported "))
        # Columns: A date | B-D Transport | E-G Food | H-J Books | K-M Zoo | N day total
        # 3 October: header row 4, two item rows (Food has two), sums row 7
        self.assertEqual((ws["A4"].value.date(), ws["A4"].number_format), (date(2026, 10, 3), "yyyy-mm-dd (ddd)"))
        self.assertEqual([ws.cell(4, c).value for c in (4, 7, 10, 13, 14)],
                         ["Transport", "Food", "Books", "Zoo", "Day total"])
        self.assertEqual(ws["G4"].fill.fgColor.rgb, "00F5C5B9")  # Food's color, lightened
        self.assertEqual([ws.cell(5, c).value for c in range(5, 8)], ["Coffee", "Espresso", 8.5])
        self.assertEqual([ws.cell(6, c).value for c in range(5, 8)], ["Restaurant", "Menu del día", 18])
        self.assertEqual((ws["B5"].value, ws["C5"].value, ws["D5"].value), ("Taxi", "To the market", 12))
        self.assertEqual((ws["A7"].value.date(), ws["D7"].value, ws["G7"].value, ws["J7"].value),
                         (date(2026, 10, 3), "=SUM(D5:D6)", "=SUM(G5:G6)", None))  # nothing in Books that day
        self.assertEqual(ws["N7"].value, "=SUM(D7,G7)")
        self.assertEqual(ws["A7"].fill.fgColor.rgb, "00FFF176")
        # 4 October: one expense per category, so rows 8-10; one without a rate shows its own amount as text
        self.assertEqual(ws["J9"].value, "US$ 5.00 (no rate)")
        self.assertEqual(ws["M9"].value, 37.5)
        self.assertEqual(ws["A10"].value.date(), date(2026, 10, 4))
        # Totals, after a blank row: the daily sums (column A only has dates on the days' first and sums rows)
        self.assertEqual((ws["A11"].value, ws["A12"].value), (None, "Total"))
        self.assertEqual(ws["G12"].value, '=SUMIF($A$4:$A$10,"<>",G4:G10)')
        self.assertEqual(ws.freeze_panes, "B1")


if __name__ == "__main__":
    unittest.main()
