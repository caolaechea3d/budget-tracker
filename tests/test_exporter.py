import csv
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from expense_tracker.catalog import default_categories
from expense_tracker.exporter import (
    EXPENSE_COLUMNS,
    INCOME_COLUMNS,
    export_expenses_csv,
    export_income_csv,
    export_xlsx,
)
from expense_tracker.importer import ImportOptions, build_expenses, guess_columns, read_csv
from expense_tracker.storage import Expense, Income

D = Decimal


def sample_expenses() -> list[Expense]:
    return [
        Expense(D("213.00"), "Food", "Plaza Vea", date(2026, 9, 29), rate=D("3.8"), subcategory="Groceries",
                payment_method="Credit card", tags=["family", "weekly"]),
        Expense(D("20"), "Shopping", "=SUM(1) is just text", date(2026, 9, 27), currency="USD", rate=D("3.75")),
        Expense(D("10"), "Transport", "Old entry\x01 without rate", date(2026, 9, 2)),
    ]


class ExportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def read_rows(self, path: Path) -> list[list[str]]:
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.reader(f))

    def test_expenses_csv(self):
        path = self.dir / "expenses.csv"
        export_expenses_csv(sample_expenses(), path)
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))  # BOM, so Excel reads accents right
        rows = self.read_rows(path)
        self.assertEqual(rows[0], [
            "Date", "Category", "Subcategory", "Description", "Payment method", "Tags",
            "Currency", "Amount", "Rate (PEN per USD)", "Amount in PEN", "Amount in USD",
        ])
        self.assertEqual(rows[1], ["2026-09-29", "Food", "Groceries", "Plaza Vea", "Credit card", "family, weekly",
                                   "PEN", "213.00", "3.80", "213.00", "56.05"])
        self.assertEqual(rows[2][6:], ["USD", "20.00", "3.75", "75.00", "20.00"])
        self.assertEqual(rows[3][8:], ["", "10.00", ""])  # no rate: no USD amount
        self.assertEqual(rows[4], [])  # blank line before the totals
        self.assertEqual(rows[5], ["Total (3; 1 without exchange rate)"] + [""] * 8 + ["298.00", "76.05"])

    def test_income_csv(self):
        path = self.dir / "income.csv"
        export_income_csv([Income(D("3000"), "Salary", "September", date(2026, 9, 30), "USD", D("3.75"))], path)
        rows = self.read_rows(path)
        self.assertEqual(rows[0], [c.title for c in INCOME_COLUMNS])
        self.assertEqual(rows[1], ["2026-09-30", "Salary", "September", "USD", "3000.00", "3.75", "11250.00", "3000.00"])
        self.assertEqual(rows[3], ["Total (1)", "", "", "", "", "", "11250.00", "3000.00"])

    def test_empty_csv_has_only_the_header(self):
        path = self.dir / "none.csv"
        export_expenses_csv([], path)
        self.assertEqual(self.read_rows(path), [[c.title for c in EXPENSE_COLUMNS]])

    def test_exported_csv_imports_back_without_errors(self):
        path = self.dir / "expenses.csv"
        export_expenses_csv(sample_expenses(), path)
        data = read_csv(path)
        cols = guess_columns(data.headers)
        self.assertEqual((cols["date"], cols["category"], cols["description"], cols["amount"]), (0, 1, 3, 7))
        opts = ImportOptions(date_col=0, amount_col=7, description_col=3, category_col=1, rate=D("3.7"))
        result = build_expenses(data, opts, [], default_categories())
        self.assertEqual((len(result.expenses), result.errors), (3, []))  # the totals line is skipped

    def test_excel_workbook(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl isn't installed")
        path = self.dir / "budget.xlsx"
        export_xlsx(path, sample_expenses(), [], "September 2026", "Ana", "Category: Food")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Expenses", "Income"])
        ws = wb["Expenses"]
        self.assertEqual(ws["A1"].value, "Expenses — September 2026")
        self.assertTrue(ws["A2"].value.startswith("Ana · Category: Food · Exported "))
        self.assertEqual([c.value for c in ws[4]], [c.title for c in EXPENSE_COLUMNS])
        self.assertEqual((ws["A5"].value.date(), ws["A5"].number_format), (date(2026, 9, 29), "d mmm yyyy"))
        self.assertEqual((ws["J5"].value, ws["J5"].number_format), (213, '"S/" #,##0.00'))
        self.assertEqual((ws["H6"].value, ws["H6"].number_format), (20, '"US$" #,##0.00'))  # in its own currency
        self.assertEqual((ws["D6"].value, ws["D6"].data_type), ("=SUM(1) is just text", "s"))
        self.assertEqual(ws["D7"].value, "Old entry without rate")  # control character dropped
        self.assertIsNone(ws["K7"].value)
        # Totals follow Excel's filter (SUBTOTAL skips filtered-out rows)
        self.assertEqual(ws["A8"].value, '="Total ("&SUBTOTAL(103,H5:H7)&")"')
        self.assertEqual((ws["J8"].value, ws["K8"].value), ("=SUBTOTAL(109,J5:J7)", "=SUBTOTAL(109,K5:K7)"))
        self.assertEqual((ws.auto_filter.ref, ws.freeze_panes), ("A4:K7", "A5"))
        income = wb["Income"]
        self.assertEqual((income["A1"].value, income["A5"].value), ("Income — September 2026", "No income"))
        self.assertTrue(income["A2"].value.startswith("Ana · Exported "))  # expense filters don't apply here


if __name__ == "__main__":
    unittest.main()
